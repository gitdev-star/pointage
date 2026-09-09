# =====================================================
# PATH: pointage/django_hr/leaves/views.py
# =====================================================

from django.utils import timezone
from datetime import date, datetime, timedelta
import django_filters
from django.http import HttpResponse
from rest_framework import viewsets, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend
from accounts.permissions import require_perm
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


from employees.models import WorkSchedule, Employee
from django.db.models import OuterRef, Subquery, CharField
from django.db.models.functions import Cast
from audit_log.models import AuditLog
from audit_log.utils import log_action, diff_dict, snapshot

from .models import LeaveType, LeaveBalance, LeaveRequest, MaternityLeave
from .serializers import (
    LeaveTypeSerializer,
    LeaveBalanceSerializer,
    LeaveRequestSerializer,
    MaternityLeaveSerializer,
)
from rest_framework.exceptions import PermissionDenied

STATUS_LABELS_FR = {
    "PENDING":   "En attente",
    "APPROVED":  "Approuvé",
    "REJECTED":  "Rejeté",
    "CANCELLED": "Annulé",
}

HOURS_PER_DAY = 8.0

BREASTFEEDING_START             = "07:30"
BREASTFEEDING_END               = "15:30"
BREASTFEEDING_EARLY_LEAVE_LIMIT = "15:27"   # même marge de 3 min que l'horaire standard (16:27 pour 16:30)
BREASTFEEDING_WORK_HOURS        = 8.0        # 07:30 → 15:30, sans déduction pause (cf. analysis_service)
BREASTFEEDING_OVERTIME_HOURS    = 8.5
BREASTFEEDING_DURATION_WEEKS    = 15

def get_hr_name(request):
    """Get HR username from request — works with both middleware and require_perm."""
    # Try hr_profile first (set by require_perm)
    profile = getattr(request, 'hr_profile', None)
    if profile:
        return profile.username
    # Fallback: use hr_username set by HRJWTMiddleware
    username = getattr(request, 'hr_username', None)
    if username:
        return username
    # Last resort: query by hr_user_id
    user_id = getattr(request, 'hr_user_id', None)
    if user_id:
        try:
            from accounts.models import HRProfile
            p = HRProfile.objects.get(auth_user_id=user_id)
            return p.username
        except Exception:
            pass
    return "Inconnu"


from alerts.email_utils import (
    notify_leave_created,
    notify_maternity_created,
    notify_maternity_ending_soon,
    notify_maternity_returned,
    notify_maternity_extended,
)




def resolve_balance_impact(leave_type, duration_hours, days_requested):
    """
    Normally a leave request debits its own leave_type's balance.
    Permissions en heure (PM) are the exception: they're recorded under PM
    for history, but the day-equivalent (hours / 8) is deducted from the
    employee's CD (congé) balance instead.
    """
    if leave_type.code == "PM":
        cd_type = LeaveType.objects.filter(code="CD").first()
        hours = duration_hours or 0
        days_amount = round(float(hours) / HOURS_PER_DAY, 1)
        return (cd_type or leave_type), days_amount
    return leave_type, days_requested



class LeaveTypeViewSet(viewsets.ModelViewSet):
    ordering = ["id"]
    queryset = LeaveType.objects.filter(is_active=True).only(
        "id", "code", "name", "days_per_year",
        "is_paid", "requires_document", "color", "is_active", "is_protected",
    )
    serializer_class = LeaveTypeSerializer

    def get_permissions(self):
        write_actions = {"create", "update", "partial_update", "destroy"}
        perm_key = "leaves_write" if self.action in write_actions else "leaves_read"
        return [require_perm(perm_key)()]

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        if instance.code in PROTECTED_LEAVE_CODES:
            raise PermissionDenied(
                f"Le type « {instance.name} » est protégé et ne peut pas être supprimé."
            )
        return super().destroy(request, *args, **kwargs)




class LeaveBalanceViewSet(viewsets.ModelViewSet):
    queryset = (
        LeaveBalance.objects
        .select_related("employee", "leave_type")
        .only(
            "id", "year", "entitled_days", "used_days", "pending_days",
            "employee__id", "employee__first_name", "employee__last_name",
            "employee__employee_id",
            "leave_type__id", "leave_type__name", "leave_type__code",
        )
    )
    serializer_class = LeaveBalanceSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ["employee", "leave_type", "year"]

    def list(self, request, *args, **kwargs):
        employee_id = request.query_params.get("employee")
        year = request.query_params.get("year")
        if employee_id and year:
            employee = Employee.objects.filter(pk=employee_id).first()
            if employee:
                existing_ids = set(
                    LeaveBalance.objects.filter(employee=employee, year=year)
                    .values_list("leave_type_id", flat=True)
                )
                for lt in LeaveType.objects.filter(is_active=True).exclude(id__in=existing_ids):
                    LeaveBalance.get_or_create_for(employee, lt, int(year))
        return super().list(request, *args, **kwargs)
    
    def get_permissions(self):
        write_actions = {"create", "update", "partial_update", "destroy"}
        perm_key = "leaves_write" if self.action in write_actions else "leaves_read"
        return [require_perm(perm_key)()]

class LeaveRequestFilter(django_filters.FilterSet):
    date_from = django_filters.DateFilter(method="filter_date_from")
    date_to   = django_filters.DateFilter(method="filter_date_to")

    class Meta:
        model  = LeaveRequest
        fields = ["employee", "leave_type", "employee__factory"]

    def filter_date_from(self, queryset, name, value):
        # Garde les événements encore actifs à partir de date_from
        # (chevauchement, pas juste ceux qui commencent après)
        return queryset.filter(end_date__gte=value)

    def filter_date_to(self, queryset, name, value):
        # Garde les événements déjà commencés avant date_to
        return queryset.filter(start_date__lte=value)

EXCEL_HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

EXCEL_WEEKEND_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF200",
)

EXCEL_EVENT_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6",
)

EXCEL_OK_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9",
)

EXCEL_TITLE_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

EXCEL_THIN_BORDER = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)

FRENCH_WEEKDAYS = [
    "lundi",
    "mardi",
    "mercredi",
    "jeudi",
    "vendredi",
    "samedi",
    "dimanche",
]

FRENCH_MONTHS = [
    "",
    "janvier",
    "février",
    "mars",
    "avril",
    "mai",
    "juin",
    "juillet",
    "août",
    "septembre",
    "octobre",
    "novembre",
    "décembre",
]


def get_export_month(request):
    """
    Récupère le mois depuis date_from.

    Exemple :
    ?date_from=2026-09-01

    Si date_from est absent ou invalide,
    utilise le mois actuel.
    """
    value = request.query_params.get("date_from")

    if value:
        try:
            selected_date = datetime.strptime(
                value,
                "%Y-%m-%d",
            ).date()

            return selected_date.year, selected_date.month

        except ValueError:
            pass

    today = timezone.localdate()

    return today.year, today.month


def format_duration_from_times(start_time, end_time):
    """
    Retourne la durée en minutes.
    """
    if not start_time or not end_time:
        return None

    reference_date = date.today()

    start_datetime = datetime.combine(
        reference_date,
        start_time,
    )

    end_datetime = datetime.combine(
        reference_date,
        end_time,
    )

    total_minutes = int(
        (
            end_datetime - start_datetime
        ).total_seconds() // 60
    )

    if total_minutes <= 0:
        return None

    return total_minutes


def get_half_day(start_time):
    if not start_time:
        return ""

    if start_time.hour < 12:
        return "MATIN"

    return "APRÈS-MIDI"


def apply_cell_style(cell, fill=None, bold=False):
    cell.border = EXCEL_THIN_BORDER
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    if fill:
        cell.fill = fill

    cell.font = Font(
        name="Calibri",
        size=10,
        bold=bold,
    )


def autosize_worksheet(worksheet, maximum_width=30):
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        maximum_length = 0

        for cell in column_cells:
            value = cell.value

            if value is None:
                continue

            value_length = len(str(value))
            maximum_length = max(
                maximum_length,
                value_length,
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 2,
            maximum_width,
        )

    
class LeaveRequestViewSet(viewsets.ModelViewSet):
    queryset = (
        LeaveRequest.objects
        .select_related("employee", "employee__factory", "leave_type")
        .only(
            "id",
            "start_date",
            "end_date",
            "start_time",
            "end_time",
            "days_requested",
            "duration_hours",
            "reason",
            "document",
            "created_at",
            "updated_at",
            "employee__id",
            "employee__first_name",
            "employee__last_name",
            "employee__employee_id",
            "employee__factory__id",
            "employee__factory__name",
            "leave_type__id",
            "leave_type__name",
            "leave_type__code",
        )
    )
    serializer_class = LeaveRequestSerializer

    filter_backends  = [DjangoFilterBackend, filters.OrderingFilter, filters.SearchFilter]
    filterset_class  = LeaveRequestFilter
    search_fields    = ["employee__first_name", "employee__last_name", "employee__employee_id"]
    ordering_fields  = ["start_date", "created_at"]
    
    def get_queryset(self):
        qs = super().get_queryset()
        latest_log = AuditLog.objects.filter(
            model_name="leaverequest",
            object_id=Cast(OuterRef("pk"), output_field=CharField()),
        ).order_by("-timestamp")
        qs = qs.annotate(
            last_action=Subquery(latest_log.values("action")[:1]),
            last_action_at=Subquery(latest_log.values("timestamp")[:1]),
            last_action_by=Subquery(latest_log.values("username")[:1]),
        )
        return qs



    def perform_create(self, serializer):
        leave = serializer.save()
        log_action(self.request, leave, "CREATE")
        try:
            name = get_hr_name(self.request)
            notify_leave_created(leave, triggered_by=name)
        except Exception as e:
            print(f"[NOTIFY] leave_created error: {e}")


    def perform_update(self, serializer):
        old_data = snapshot(serializer.instance)
        leave = serializer.save()
        new_data = snapshot(leave)
        changes = diff_dict(old_data, new_data)
        if changes:
            log_action(self.request, leave, "UPDATE", changes)


    def perform_destroy(self, instance):
        log_action(self.request, instance, "DELETE")
        super().perform_destroy(instance)
            

    @action(detail=False, methods=["get"])
    def export(self, request):
        year, month = get_export_month(request)

        first_day = date(year, month, 1)

        last_day = date(
            year,
            month,
            monthrange(year, month)[1],
        )

        number_of_days = last_day.day

        # ---------------------------------------------
        # Employés concernés
        # ---------------------------------------------

        employees = (
            Employee.objects
            .filter(status=Employee.Status.ACTIVE)
            .select_related(
                "factory",
                "department",
                "section",
            )
            .order_by(
                "factory__name",
                "section__name",
                "last_name",
                "first_name",
            )
        )

        factory_id = request.query_params.get(
            "employee__factory"
        )

        if factory_id:
            employees = employees.filter(
                factory_id=factory_id
            )

        search = request.query_params.get(
            "search",
            ""
        ).strip()

        if search:
            from django.db.models import Q

            employees = employees.filter(
                Q(first_name__icontains=search)
                | Q(last_name__icontains=search)
                | Q(employee_id__icontains=search)
                | Q(matricule_paie__icontains=search)
            )

        employees = list(employees)

        employee_ids = [
            employee.id
            for employee in employees
        ]

        # ---------------------------------------------
        # Événements qui chevauchent le mois
        # ---------------------------------------------

        events = (
            LeaveRequest.objects
            .filter(
                employee_id__in=employee_ids,
                start_date__lte=last_day,
                end_date__gte=first_day,
            )
            .select_related(
                "employee",
                "employee__factory",
                "employee__department",
                "employee__section",
                "leave_type",
            )
            .order_by(
                "employee__last_name",
                "employee__first_name",
                "start_date",
                "start_time",
            )
        )

        leave_type_id = request.query_params.get(
            "leave_type"
        )

        if leave_type_id:
            events = events.filter(
                leave_type_id=leave_type_id
            )

        events = list(events)

        # ---------------------------------------------
        # Regrouper les événements par employé et date
        # ---------------------------------------------

        events_by_employee_and_day = {}

        events_by_employee = {}

        for event in events:
            events_by_employee.setdefault(
                event.employee_id,
                [],
            ).append(event)

            current_date = max(
                event.start_date,
                first_day,
            )

            event_last_date = min(
                event.end_date,
                last_day,
            )

            while current_date <= event_last_date:
                key = (
                    event.employee_id,
                    current_date,
                )

                codes = events_by_employee_and_day.setdefault(
                    key,
                    [],
                )

                if event.leave_type.code not in codes:
                    codes.append(event.leave_type.code)

                current_date += timedelta(days=1)

        # ---------------------------------------------
        # Codes à afficher dans le résumé
        # ---------------------------------------------

        leave_codes = list(
            LeaveType.objects
            .filter(is_active=True)
            .order_by("id")
            .values_list("code", flat=True)
        )

        # ---------------------------------------------
        # Création du classeur
        # ---------------------------------------------

        workbook = Workbook()

        monthly_sheet = workbook.active
        monthly_sheet.title = "Suivi mensuel"

        detail_sheet = workbook.create_sheet(
            title="Détails événements"
        )

        # =============================================
        # FEUILLE 1 : SUIVI MENSUEL
        # =============================================

        fixed_headers = [
            "Matricule",
            "Matricule paie",
            "Nom",
            "Prénom",
            "Site",
            "Section",
            "Date d’entrée",
        ]

        total_columns = (
            len(fixed_headers)
            + number_of_days
            + len(leave_codes)
            + 1
        )

        monthly_sheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=total_columns,
        )

        title_cell = monthly_sheet.cell(
            row=1,
            column=1,
            value=(
                f"Suivi mensuel des événements — "
                f"{FRENCH_MONTHS[month].capitalize()} {year}"
            ),
        )

        title_cell.fill = EXCEL_TITLE_FILL
        title_cell.font = Font(
            color="FFFFFF",
            bold=True,
            size=14,
        )
        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        monthly_sheet.row_dimensions[1].height = 24

        # En-têtes fixes
        for column_index, header in enumerate(
            fixed_headers,
            start=1,
        ):
            cell = monthly_sheet.cell(
                row=2,
                column=column_index,
                value=header,
            )

            apply_cell_style(
                cell,
                fill=EXCEL_HEADER_FILL,
                bold=True,
            )

        # En-têtes des jours
        first_day_column = len(fixed_headers) + 1

        for day_number in range(1, number_of_days + 1):
            current_date = date(
                year,
                month,
                day_number,
            )

            column_index = (
                first_day_column
                + day_number
                - 1
            )

            weekday_name = FRENCH_WEEKDAYS[
                current_date.weekday()
            ]

            cell = monthly_sheet.cell(
                row=2,
                column=column_index,
                value=(
                    f"{weekday_name}\n"
                    f"{day_number:02d}/"
                    f"{month:02d}"
                ),
            )

            fill = (
                EXCEL_WEEKEND_FILL
                if current_date.weekday() >= 5
                else EXCEL_HEADER_FILL
            )

            apply_cell_style(
                cell,
                fill=fill,
                bold=True,
            )

        # En-têtes des totaux par code
        first_summary_column = (
            first_day_column
            + number_of_days
        )

        for offset, code in enumerate(
            leave_codes
        ):
            cell = monthly_sheet.cell(
                row=2,
                column=first_summary_column + offset,
                value=code,
            )

            apply_cell_style(
                cell,
                fill=EXCEL_HEADER_FILL,
                bold=True,
            )

        ok_summary_column = (
            first_summary_column
            + len(leave_codes)
        )

        ok_header = monthly_sheet.cell(
            row=2,
            column=ok_summary_column,
            value="OK",
        )

        apply_cell_style(
            ok_header,
            fill=EXCEL_OK_FILL,
            bold=True,
        )

        # Données des employés
        for row_index, employee in enumerate(
            employees,
            start=3,
        ):
            fixed_values = [
                employee.employee_id or "",
                employee.matricule_paie or "",
                employee.last_name or "",
                employee.first_name or "",
                (
                    employee.factory.name
                    if employee.factory
                    else ""
                ),
                (
                    employee.section.name
                    if employee.section
                    else ""
                ),
                employee.hire_date,
            ]

            for column_index, value in enumerate(
                fixed_values,
                start=1,
            ):
                cell = monthly_sheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value,
                )

                apply_cell_style(cell)

            hire_date_cell = monthly_sheet.cell(
                row=row_index,
                column=7,
            )

            if employee.hire_date:
                hire_date_cell.number_format = "dd/mm/yyyy"

            employee_code_counts = {
                code: 0
                for code in leave_codes
            }

            ok_count = 0

            for day_number in range(
                1,
                number_of_days + 1,
            ):
                current_date = date(
                    year,
                    month,
                    day_number,
                )

                column_index = (
                    first_day_column
                    + day_number
                    - 1
                )

                codes = events_by_employee_and_day.get(
                    (
                        employee.id,
                        current_date,
                    ),
                    [],
                )

                if codes:
                    displayed_value = " / ".join(codes)
                    fill = EXCEL_EVENT_FILL

                    for code in codes:
                        if code in employee_code_counts:
                            employee_code_counts[code] += 1
                else:
                    displayed_value = "OK"
                    fill = (
                        EXCEL_WEEKEND_FILL
                        if current_date.weekday() >= 5
                        else EXCEL_OK_FILL
                    )
                    ok_count += 1

                cell = monthly_sheet.cell(
                    row=row_index,
                    column=column_index,
                    value=displayed_value,
                )

                apply_cell_style(
                    cell,
                    fill=fill,
                    bold=bool(codes),
                )

            # Totaux par code
            for offset, code in enumerate(
                leave_codes
            ):
                count = employee_code_counts[code]

                cell = monthly_sheet.cell(
                    row=row_index,
                    column=first_summary_column + offset,
                    value=count if count else "-",
                )

                apply_cell_style(cell)

            ok_cell = monthly_sheet.cell(
                row=row_index,
                column=ok_summary_column,
                value=ok_count,
            )

            apply_cell_style(
                ok_cell,
                fill=EXCEL_OK_FILL,
                bold=True,
            )

        monthly_sheet.freeze_panes = "H3"
        monthly_sheet.auto_filter.ref = (
            f"A2:{get_column_letter(total_columns)}"
            f"{len(employees) + 2}"
        )

        monthly_sheet.row_dimensions[2].height = 44

        monthly_sheet.column_dimensions["A"].width = 14
        monthly_sheet.column_dimensions["B"].width = 16
        monthly_sheet.column_dimensions["C"].width = 22
        monthly_sheet.column_dimensions["D"].width = 22
        monthly_sheet.column_dimensions["E"].width = 14
        monthly_sheet.column_dimensions["F"].width = 20
        monthly_sheet.column_dimensions["G"].width = 14

        for column_index in range(
            first_day_column,
            first_summary_column,
        ):
            monthly_sheet.column_dimensions[
                get_column_letter(column_index)
            ].width = 11

        for column_index in range(
            first_summary_column,
            total_columns + 1,
        ):
            monthly_sheet.column_dimensions[
                get_column_letter(column_index)
            ].width = 12

        # =============================================
        # FEUILLE 2 : DÉTAILS DES ÉVÉNEMENTS
        # =============================================

        detail_headers = [
            "Matricule",
            "Matricule paie",
            "Nom",
            "Prénom",
            "Date départ",
            "Date retour",
            "Matin/Après-midi",
            "Heure départ",
            "Heure retour",
            "Libellé",
            "Code",
            "Nombre d’heures",
            "Nombre de jours",
            "Site",
            "Section",
            "Motif",
        ]

        for column_index, header in enumerate(
            detail_headers,
            start=1,
        ):
            cell = detail_sheet.cell(
                row=1,
                column=column_index,
                value=header,
            )

            apply_cell_style(
                cell,
                fill=EXCEL_HEADER_FILL,
                bold=True,
            )

        for row_index, event in enumerate(
            events,
            start=2,
        ):
            employee = event.employee

            duration_minutes = format_duration_from_times(
                event.start_time,
                event.end_time,
            )

            row_values = [
                employee.employee_id or "",
                employee.matricule_paie or "",
                employee.last_name or "",
                employee.first_name or "",
                event.start_date,
                event.end_date,
                get_half_day(event.start_time),
                event.start_time,
                event.end_time,
                event.leave_type.name,
                event.leave_type.code,
                (
                    duration_minutes / 1440
                    if duration_minutes is not None
                    else None
                ),
                event.days_requested,
                (
                    employee.factory.name
                    if employee.factory
                    else ""
                ),
                (
                    employee.section.name
                    if employee.section
                    else ""
                ),
                event.reason or "",
            ]

            for column_index, value in enumerate(
                row_values,
                start=1,
            ):
                cell = detail_sheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value,
                )

                apply_cell_style(cell)

            detail_sheet.cell(
                row=row_index,
                column=5,
            ).number_format = "dd/mm/yyyy"

            detail_sheet.cell(
                row=row_index,
                column=6,
            ).number_format = "dd/mm/yyyy"

            detail_sheet.cell(
                row=row_index,
                column=8,
            ).number_format = "hh:mm"

            detail_sheet.cell(
                row=row_index,
                column=9,
            ).number_format = "hh:mm"

            # Affichage réel en heures et minutes :
            # 03:37 au lieu de 3.62
            detail_sheet.cell(
                row=row_index,
                column=12,
            ).number_format = "[h]:mm"

            detail_sheet.cell(
                row=row_index,
                column=13,
            ).number_format = "0.0"

        detail_sheet.freeze_panes = "A2"

        if events:
            detail_sheet.auto_filter.ref = (
                f"A1:P{len(events) + 1}"
            )

        autosize_worksheet(
            detail_sheet,
            maximum_width=28,
        )

        # ---------------------------------------------
        # Génération de la réponse Excel
        # ---------------------------------------------

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        filename = (
            f"evenements_"
            f"{year}_"
            f"{month:02d}.xlsx"
        )

        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        response["Content-Disposition"] = (
            f'attachment; filename="{filename}"'
        )

        return response

class MaternityLeaveViewSet(viewsets.ModelViewSet):
    queryset = MaternityLeave.objects.select_related(
        "employee", "employee__factory", "employee__department"
    ).all()
    serializer_class = MaternityLeaveSerializer
    filter_backends  = [DjangoFilterBackend, filters.OrderingFilter]
    filterset_fields = ["employee", "status"]
    ordering_fields  = ["leave_start_date", "expected_birth_date"]

    def perform_create(self, serializer):
        maternity = serializer.save()
        try:
            name = get_hr_name(self.request)
            notify_maternity_created(maternity, triggered_by=name)
        except Exception as e:
            print(f"[NOTIFY] maternity_created error: {e}")


    @action(detail=True, methods=["post"])
    def mark_returned(self, request, pk=None):
        ml = self.get_object()
        ml.actual_return_date = request.data.get("return_date") or timezone.now().date()
        ml.status = MaternityLeave.Status.RETURNED
        ml.save()

        # ── Assignation automatique de l'horaire d'allaitement ────────────────
        try:
            return_date = ml.actual_return_date
            valid_until = return_date + timedelta(weeks=BREASTFEEDING_DURATION_WEEKS)

            # Si un horaire d'allaitement existe déjà pour cette employée
            # (ex: mark_returned rappelé par erreur), on le met à jour au lieu
            # d'en créer un doublon (name a unique=True sur WorkSchedule).
            existing = WorkSchedule.objects.filter(
                employee=ml.employee,
                name__startswith="Allaitement",
            ).first()

            if existing:
                existing.standard_start     = BREASTFEEDING_START
                existing.standard_end       = BREASTFEEDING_END
                existing.early_leave_limit  = BREASTFEEDING_EARLY_LEAVE_LIMIT
                existing.standard_work_hours = BREASTFEEDING_WORK_HOURS
                existing.overtime_threshold_hours = BREASTFEEDING_OVERTIME_HOURS
                existing.valid_from  = return_date
                existing.valid_until = valid_until
                existing.is_active   = True
                existing.description = "Horaire d'allaitement — assigné automatiquement après congé maternité"
                existing.save()
            else:
                WorkSchedule.objects.create(
                    employee=ml.employee,
                    name=f"Allaitement - {ml.employee.employee_id}",
                    description="Horaire d'allaitement — assigné automatiquement après congé maternité",
                    standard_start=BREASTFEEDING_START,
                    standard_end=BREASTFEEDING_END,
                    early_leave_limit=BREASTFEEDING_EARLY_LEAVE_LIMIT,
                    standard_work_hours=BREASTFEEDING_WORK_HOURS,
                    overtime_threshold_hours=BREASTFEEDING_OVERTIME_HOURS,
                    valid_from=return_date,
                    valid_until=valid_until,
                    is_active=True,
                )
        except Exception as e:
            print(f"[WORKSCHEDULE] auto-assign breastfeeding schedule error: {e}")

        try:
            name = get_hr_name(request)
            notify_maternity_returned(ml, triggered_by=name)
        except Exception as e:
            print(f"[NOTIFY] maternity_returned error: {e}")

        return Response(MaternityLeaveSerializer(ml).data)

    @action(detail=True, methods=["post"])
    def extend(self, request, pk=None):
        ml = self.get_object()
        extended_end = request.data.get("extended_end_date")
        if not extended_end:
            return Response({"detail": "extended_end_date requis."}, status=400)
        ml.extended_end_date = extended_end
        ml.status = MaternityLeave.Status.EXTENDED
        ml.note   = request.data.get("note", ml.note)
        ml.save()
        try:
            name = get_hr_name(request)
            notify_maternity_extended(ml, triggered_by=name)
        except Exception as e:
            print(f"[NOTIFY] maternity_extended error: {e}")
        return Response(MaternityLeaveSerializer(ml).data)

    @action(detail=False, methods=["get"])
    def active(self, request):
        today = timezone.now().date()
        qs = self.queryset.filter(
            status__in=["DECLARED", "ON_LEAVE", "EXTENDED"],
            leave_end_date__gte=today,
        )
        return Response(MaternityLeaveSerializer(qs, many=True).data)

    @action(detail=False, methods=["get"])
    def upcoming_returns(self, request):
        today     = timezone.now().date()
        threshold = today + timedelta(days=30)
        qs = self.queryset.filter(
            status__in=["ON_LEAVE", "EXTENDED"],
            leave_end_date__lte=threshold,
            leave_end_date__gte=today,
            actual_return_date__isnull=True,
        )
        return Response(MaternityLeaveSerializer(qs, many=True).data)

    @action(detail=False, methods=["get"])
    def ending_soon(self, request):
        """Maternity leaves ending in <= 7 days — also sends alert emails."""
        today     = timezone.now().date()
        threshold = today + timedelta(days=7)
        qs = self.queryset.filter(
            status__in=["ON_LEAVE", "EXTENDED"],
            leave_end_date__lte=threshold,
            leave_end_date__gte=today,
        )
        for ml in qs:
            days_remaining = (ml.leave_end_date - today).days
            try:
                notify_maternity_ending_soon(ml, days_remaining=days_remaining)
            except Exception as e:
                print(f"[NOTIFY] maternity_ending_soon error: {e}")
        return Response(MaternityLeaveSerializer(qs, many=True).data)

    def get_permissions(self):
        write_actions = {"create", "update", "partial_update", "destroy", "mark_returned", "extend"}
        perm_key = "leaves_write" if self.action in write_actions else "leaves_read"
        return [require_perm(perm_key)()]