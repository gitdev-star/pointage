from io import BytesIO
from pathlib import Path

from django.utils import timezone
from docx import Document


MODELE_FICHE_POSTE = (
    Path(__file__).resolve().parent.parent
    / "templates"
    / "fiche_de_poste.docx"
)


def _ajouter_valeur(paragraphe, valeur):
    valeur = str(valeur or "").strip()
    if valeur:
        paragraphe.add_run(f"\n{valeur}")


def _remplacer_texte(paragraphe, ancien, nouveau):
    if ancien not in paragraphe.text:
        return

    texte = paragraphe.text.replace(ancien, nouveau)
    if paragraphe.runs:
        paragraphe.runs[0].text = texte
        for run in paragraphe.runs[1:]:
            run.text = ""
    else:
        paragraphe.add_run(texte)


def _paragraphes_document(document):
    yield from document.paragraphs

    for table in document.tables:
        cellules_vues = set()
        for ligne in table.rows:
            for cellule in ligne.cells:
                identifiant = id(cellule._tc)
                if identifiant in cellules_vues:
                    continue
                cellules_vues.add(identifiant)
                yield from cellule.paragraphs


def _responsable_hierarchique(demande):
    manager = getattr(demande.departement, "manager", None)
    if not manager:
        return ""

    nom = manager.full_name.strip()
    poste = getattr(manager, "job_title", None)
    if poste:
        return f"{nom} - {poste}"
    return nom


def generer_fiche_poste(embauche):
    """Génère la fiche de poste Word préremplie d'un candidat."""
    candidat = embauche.candidat
    demande = candidat.processus.demande
    document = Document(MODELE_FICHE_POSTE)

    manager = _responsable_hierarchique(demande)
    poste_description = (demande.poste.description or "").strip()
    formation = "\n".join(
        partie
        for partie in [
            (demande.profil_diplome or "").strip(),
            (demande.experience_professionnelle or "").strip(),
        ]
        if partie
    )
    competences = "\n".join(
        partie
        for partie in [
            (demande.competences_techniques or "").strip(),
            (demande.savoir_faire or "").strip(),
        ]
        if partie
    )

    table = document.tables[1]
    table.rows[0].cells[1].text = demande.poste.name
    table.rows[1].cells[1].text = manager
    table.rows[2].cells[1].text = ""
    table.rows[3].cells[1].text = demande.departement.name

    # Les cellules des cinq dernières lignes sont fusionnées. On cible
    # leurs paragraphes par position afin de conserver la mise en page
    # du modèle, indépendamment de l'encodage de ses libellés.
    _ajouter_valeur(table.rows[4].cells[0].paragraphs[0], poste_description)
    _ajouter_valeur(table.rows[5].cells[0].paragraphs[0], poste_description)
    _ajouter_valeur(
        table.rows[6].cells[0].paragraphs[0],
        demande.designation_taches,
    )
    _ajouter_valeur(table.rows[7].cells[0].paragraphs[0], formation)
    _ajouter_valeur(table.rows[8].cells[0].paragraphs[3], competences)
    _ajouter_valeur(
        table.rows[8].cells[0].paragraphs[6],
        demande.savoir_etre,
    )

    date_generation = timezone.localdate().strftime("%d/%m/%Y")
    matricule = ""
    if embauche.employe_id and embauche.employe:
        matricule = embauche.employe.employee_id

    for paragraphe in _paragraphes_document(document):
        _remplacer_texte(
            paragraphe,
            "date de generation",
            date_generation,
        )

    identite = document.paragraphs[1]
    identite.text = (
        f"Nom et Prénoms de l’Employé(e) : {candidat.nom_complet}\n"
        f"N° IML : {matricule}\n"
        f"Date : {date_generation}"
    )

    contenu = BytesIO()
    document.save(contenu)
    contenu.seek(0)
    return contenu
