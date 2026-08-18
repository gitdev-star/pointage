# =====================================================
# PATH: pointage/django_hr/documents/views.py
# Fills real .docx templates with employee data
# =====================================================
import logging
import os
import re
import io
import random
import zipfile
from datetime import date, datetime, timedelta
from dateutil.relativedelta import relativedelta
from django.http import HttpResponse
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response
from employees.models import Employee
from accounts.permissions import IsHRUser
import copy 
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from .badge import build_badge_pdf

logger = logging.getLogger(__name__)

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "templates")

# ── helpers ───────────────────────────────────────────────────────────────────

MONTHS_FR = ["","janvier","f\u00e9vrier","mars","avril","mai","juin",
             "juillet","ao\u00fbt","septembre","octobre","novembre","d\u00e9cembre"]


def fmt_date(value):
    if not value:
        return "\u2014"
    if isinstance(value, str):
        try:
            value = datetime.strptime(value[:10], "%Y-%m-%d").date()
        except ValueError:
            return value
    return f"{value.day:02d} {MONTHS_FR[value.month]} {value.year}"

def add_months(date_val, n):
    if isinstance(date_val, str):
        date_val = datetime.strptime(date_val[:10], "%Y-%m-%d").date()
    return fmt_date(date_val + relativedelta(months=n))

def civilite(emp):
    return "Madame" if (getattr(emp,"sexe","") or "").upper().startswith("F") else "Monsieur"

def today_fr():
    return fmt_date(date.today())

def ref_rh():
    return f"RH-{date.today().year}/{random.randint(1,99999):05d}"

def get_classif(emp):
    return getattr(emp,"classification",None) or getattr(emp,"category",None) or "\u2014"

# ── core replace (handles split runs) ────────────────────────────────────────

def replace_in_paragraph(para, reps):
    full = "".join(r.text for r in para.runs)
    new  = full
    for old, val in reps.items():
        new = new.replace(old, str(val) if val is not None else "\u2014")
    if new == full:
        return
    if para.runs:
        para.runs[0].text = new
        for r in para.runs[1:]:
            r.text = ""

def remove_underline(doc):
    """Remove underline formatting from all runs in the document."""
    for p in doc.paragraphs:
        for r in p.runs:
            if r.underline:
                r.underline = False
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.underline:
                            r.underline = False

def replace_mergefields(doc, field_map):
    """Remplace les champs MERGEFIELD Word par du texte statique."""
    root = doc.element
    paragraphs = root.findall('.//' + qn('w:p'))
    for p in paragraphs:
        runs = list(p.findall(qn('w:r')))
        i = 0
        while i < len(runs):
            r = runs[i]
            fld = r.find(qn('w:fldChar'))
            if fld is not None and fld.get(qn('w:fldCharType')) == 'begin':
                collected = [r]
                field_name = None
                j = i + 1
                while j < len(runs):
                    rj = runs[j]
                    collected.append(rj)
                    instr = rj.find(qn('w:instrText'))
                    if instr is not None and instr.text and 'MERGEFIELD' in instr.text:
                        m = re.search(r'MERGEFIELD\s+"?([^"\s]+)"?', instr.text)
                        if m:
                            field_name = m.group(1)
                    fld2 = rj.find(qn('w:fldChar'))
                    if fld2 is not None and fld2.get(qn('w:fldCharType')) == 'separate':
                        j += 1
                        break
                    j += 1
                template_run = None
                k = j
                while k < len(runs):
                    rk = runs[k]
                    collected.append(rk)
                    if template_run is None and rk.find(qn('w:t')) is not None:
                        template_run = rk
                    fldend = rk.find(qn('w:fldChar'))
                    if fldend is not None and fldend.get(qn('w:fldCharType')) == 'end':
                        break
                    k += 1

                value = field_map.get(field_name, "") if field_name else ""

                new_r = OxmlElement('w:r')
                if template_run is not None:
                    rpr = template_run.find(qn('w:rPr'))
                    if rpr is not None:
                        new_r.append(copy.deepcopy(rpr))
                t = OxmlElement('w:t')
                t.set(qn('xml:space'), 'preserve')
                t.text = str(value)
                new_r.append(t)

                collected[0].addprevious(new_r)
                for el in collected:
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)

                runs = list(p.findall(qn('w:r')))
                i = runs.index(new_r) + 1
                continue
            i += 1

def fill_template(tpl_path, reps, strip_underline=False, mergefields=None):
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(tpl_path)

    if mergefields:
        replace_mergefields(doc, mergefields)

    for p in doc.paragraphs:
        replace_in_paragraph(p, reps)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, reps)
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            replace_in_paragraph(p, reps)
        for p in sec.footer.paragraphs:
            replace_in_paragraph(p, reps)
    if strip_underline:
        remove_underline(doc)
    # Remove trailing blank paragraphs that cause extra blank pages
    body = doc.element.body
    children = list(body)
    for child in reversed(children):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'sectPr':
            continue
        if tag == 'p':
            text = ''.join(n.text or '' for n in child.iter() if hasattr(n, 'text') and n.tag.split('}')[-1] == 't')
            has_break = child.find('.//' + qn('w:br')) is not None
            if not text.strip() and not has_break:
                body.remove(child)
                continue
        break
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()

# ── builders ──────────────────────────────────────────────────────────────────

def build_attestation(emp, extra):
    usage        = extra.get("usage") or "Pr\u00eat bancaire BOA Tanjombato"
    ref          = extra.get("ref")   or ref_rh()
    today        = today_fr()
    hire         = fmt_date(emp.hire_date)
    civ          = civilite(emp)
    name         = f"{emp.last_name} {emp.first_name}"
    classif      = get_classif(emp)
    factory      = emp.factory.name if emp.factory else "\u2014"
    contrat_txt  = "ind\u00e9termin\u00e9e" if (emp.contract_type or "") == "CDI" else "d\u00e9termin\u00e9e"
    return fill_template(os.path.join(TEMPLATES_DIR,"BASE_ATTESTATION.docx"), {
        "Date demande attestation"                                    : today,
        "HR-2026/00418"                                               : ref,
        "Pr\u00eat bancaire Boa Tanjombato"                          : usage,
        "Civilit\u00e9"                                              : civ,
        "Nom et Pr\u00e9nom"                                         : name,
        "Fonction"                                                    : emp.job_title or "\u2014",
        "Classification"                                              : classif,
        "Matricule"                                                   : emp.employee_id or "\u2014",
        "Date d\u2019embauhce"                                       : hire,
        "D\u00e9termin\u00e9e / ind\u00e9termin\u00e9e (rayer la mention inutile)" : contrat_txt,
        "Date de demande d\u2019attestation"                         : today,
        "Lieu de travail"                                             : factory,
    })

def build_certificat(emp, extra):
    ref      = extra.get("ref")      or ref_rh()
    date_fin = fmt_date(extra.get("date_fin") or date.today())
    hire     = fmt_date(emp.hire_date)
    civ      = civilite(emp)
    name     = f"{emp.last_name} {emp.first_name}"
    classif  = get_classif(emp)
    return fill_template(os.path.join(TEMPLATES_DIR,"Certificat de travail - Copie.docx"), {
        "R\u00e9f\u00e9rence RH"  : ref,
        "Civilit\u00e9\xa0"       : civ + "\xa0",
        "Civilit\u00e9"            : civ,
        "Nom et Pr\u00e9nom"       : name,
        "Emploie occup\u00e9"      : emp.job_title or "\u2014",
        "Cat\u00e9gorie pro"       : classif,
        "Matricule"                 : emp.employee_id or "\u2014",
        "Date d\u2019embauhce"     : hire,
        "Date du d\u00e9bauche"    : date_fin,
    })


def build_cdd_6(emp, extra):
    # Auto-fill salary from classification if not provided
    classif_salaire = None
    if emp.classification and emp.classification.salaire:
        try:
            val = float(emp.classification.salaire)
            if val > 0:
                classif_salaire = f"{val:,.2f}".replace(",", " ")
        except (ValueError, TypeError):
            pass
    salaire = extra.get("salaire") or emp.salaire or classif_salaire or "—"
    raw = extra.get("date_embauche") or str(emp.hire_date or date.today())

    emb_fr = fmt_date(raw)
    fin6 = add_months(raw, 6)
    essai3 = add_months(raw, 3)

    civ = civilite(emp)
    name = f"{emp.last_name} {emp.first_name}"

    fonction = emp.job_title.name if emp.job_title else "—"
    classif = emp.classification.classe if emp.classification else "—"

    b_date = fmt_date(getattr(emp, "birth_date", None))
    b_place = getattr(emp, "birth_place", None) or "—"
    addr = getattr(emp, "address", None) or "—"
    cin = getattr(emp, "cin", None) or "—"
    cin_date = fmt_date(getattr(emp, "cin_date", None))
    cin_plc = getattr(emp, "cin_place", None) or "—"

    return fill_template(
        os.path.join(TEMPLATES_DIR, "cdd-6mois.docx"),
        {
            # Ancien template avec placeholders
            "Madame/Monsieur": civ,
            "Nom et Prénom": name,
            "Date de naissance": b_date,
            "Date de naissace": b_date,
            "Lieu de naissance": b_place,
            "Adresse": addr,
            "Numéro CIN": cin,
            "Date CIN": cin_date,
            "Lieu CIN": cin_plc,
            "Emploie occupé": fonction,
            "Fonction": fonction,
            "Numéro Matricule": emp.employee_id or "—",
            "Matricule": emp.employee_id or "—",
            "Montant": salaire,
            "Classification": classif,
            "Date d’embacuhe  + mois": fin6,
            "Date d’embauche + 6 mois": fin6,
            "Date d’embauche + 3 mois": essai3,
            "Date d’embauhce + 3 mois": essai3,
            "Date d’embauhce": emb_fr,
            "Date d’embauche": emb_fr,

            # Nouveau template déjà rempli avec RAZANAMALALA Sandra
            "RAZANAMALALA Sandra": name,
            "006278": emp.employee_id or "—",
            "Machiniste": fonction,
            "309500,00": salaire,
            "309500": salaire,
            "OS1": classif,
            "04 juin 2026": emb_fr,
            "04 décembre 2026": fin6,
            "04 septembre 2026": essai3,
            "117032017870": cin,
            "19 mars 1994": b_date,
            "Sakambahiny Bemasoandro": b_place,
            "Lot IT U 41 bis - Andranonahoatra": addr,
            "29 juillet 2013": cin_date,
            ", à  Andranonahoatra.": f", à  {cin_plc}.",
            ", tao Andranonahoatra.": f", tao {cin_plc}.",
        },
        strip_underline=True,
    )

def build_evaluation_cdd(emp, extra):
    dept     = emp.department.name if emp.department else "\u2014"
    section  = extra.get("section") or dept
    factory  = emp.factory.name if emp.factory else "\u2014"
    hire     = fmt_date(emp.hire_date)
    d_debut  = fmt_date(extra.get("date_debut") or str(emp.hire_date or date.today()))
    d_fin    = fmt_date(extra.get("date_fin_eval")) if extra.get("date_fin_eval") else "\u2014"
    name     = f"{emp.last_name} {emp.first_name}"
    civ      = civilite(emp)
    today    = today_fr()
    return fill_template(os.path.join(TEMPLATES_DIR,"contrat_de_travail.docx"), {
        # Header block
        "NOM ET PRENOMS\xa0: "                          : f"NOM ET PRENOMS\xa0: {name}",
        "MATRICULE\xa0: "                               : f"MATRICULE\xa0: {emp.employee_id or chr(8212)}",
        "FONCTION\xa0:\t"                              : f"FONCTION\xa0: {emp.job_title or chr(8212)}",
        "SECTION\xa0:\t\t\t\t\t\tSite\xa0: "    : f"SECTION\xa0: {section}      Site\xa0: {factory}",
        "Date d\u2019embauche\xa0: "                   : f"Date d\u2019embauche\xa0: {hire}",
        "Date de d\u00e9but :\t\t\tDate fin : "      : f"Date de d\u00e9but : {d_debut}      Date fin : {d_fin}",
        # Signature block 1
        "Antananarivo le xxxxxxxxx"                      : f"Antananarivo le {today}",
        "Mr / Mme\xa0:  "                               : f"Mr / Mme\xa0: {civ} {name}",
        "Fonction\xa0: "                                : f"Fonction\xa0: {emp.job_title or chr(8212)}",
        "Matricule\xa0: "                               : f"Matricule\xa0: {emp.employee_id or chr(8212)}",
        "Section\xa0:\t\t\t\t\t\tSite\xa0: "     : f"Section\xa0: {section}      Site\xa0: {factory}",
        # Renewal letter date placeholders (copy 1)
        "allant du xxxxx au xxxxx"                       : f"allant du {d_debut} au {d_fin}",
        "manomboka ny xxxxxx  hatramin\u2019ny xxxxx"   : f"manomboka ny {d_debut} hatramin\u2019ny {d_fin}",
        # Renewal letter date placeholders (copy 2)
        "allant du xxxxxx au xxxxxxxxxxxxx"              : f"allant du {d_debut} au {d_fin}",
        "manomboka ny xxxxxxxxx  hatramin\u2019ny xxxxxxxxxxxxxxx" : f"manomboka ny {d_debut} hatramin\u2019ny {d_fin}",
    })

def build_cdd_18(emp, extra):
    dept = emp.department.name if emp.department else "—"
    section = extra.get("section") or (emp.section.name if emp.section else dept)
    factory = emp.factory.name if emp.factory else "—"

    # Date d'embauche
    hire_date = emp.hire_date or date.today()

    # Contrat de 18 mois
    contrat_fin = hire_date + relativedelta(months=18)

    # Début du CDI = lendemain de la fin du CDD
    debut_cdi = contrat_fin + timedelta(days=1)

    hire = fmt_date(hire_date)
    d_debut = fmt_date(hire_date)
    d_fin = fmt_date(contrat_fin)
    d_cdi = fmt_date(debut_cdi)

    name = f"{emp.last_name} {emp.first_name}"
    fonction = emp.job_title.name if emp.job_title else "—"
    matricule = emp.employee_id or "—"
    return fill_template(
        os.path.join(TEMPLATES_DIR, "cdd-18mois.docx"),
        {
            # Header block
            "ONJANIAINA Virginie"                          : name,
            "004994"                                       : matricule,
            "Machiniste"                                   : fonction,
            "GILLET ARABIE 2					Site : PBI3"    : f"{section}					Site : {factory}",
            "31 janvier 2025			Date fin :  31 juillet 2026 " : f"{d_debut}			Date fin :  {d_fin} ",
            "Date d’embauche : 31 janvier 2025"   : f"Date d’embauche : {hire}",
            # Signature + letter body (appears twice, both replaced)
            "Antananarivo le  31 juillet 2026"             : f"Antananarivo le  {d_fin}",
            "M. / Mme : ONJANIAINA Virginie"            : f"M. / Mme : {name}",
            "Matricule n° :  004994"               : f"Matricule n° :  {matricule}",
            "Fonction :  Machiniste"                    : f"Fonction :  {fonction}",
            "Section :   GILLET ARABIE 2						Site :  PBI3" : f"Section :   {section}						Site :  {factory}",
            # CDI date in letter body (appears twice)
            "à compter du 31 juillet 2026  . "          : f"à compter du {d_cdi}  . ",
            "manomboka ny 31 juillet 2026  ."              : f"manomboka ny {d_cdi}  .",
        }
    )

def build_cdd_3(emp, extra):
    raw = extra.get("date_embauche") or str(emp.hire_date or date.today())
    hire_fr = fmt_date(raw)
    fin_essai = add_months(raw, 3)

    fonction = emp.job_title.name if emp.job_title else "—"
    section = extra.get("section") or (emp.section.name if emp.section else (emp.department.name if emp.department else "—"))
    factory = emp.factory.name if emp.factory else "—"

    mergefields = {
        "Nom": emp.last_name or "—",
        "Prénom": emp.first_name or "—",
        "Matricule": emp.employee_id or "—",
        "Fonction": fonction,
        "Section": section,
        "Site_Publi": factory,
        "DE_en_LETTRES": hire_fr,
        "D_Fin_Ess_en_lettres": fin_essai,
    }
    return fill_template(
        os.path.join(TEMPLATES_DIR, "cdd-3mois.docx"),
        {},  # pas de remplacement texte classique ici
        mergefields=mergefields,
    )


def build_cdd_12(emp, extra):
    raw = extra.get("date_embauche") or str(emp.hire_date or date.today())
    hire_fr = fmt_date(raw)
    fin6 = add_months(raw, 6)
    fin12 = add_months(raw, 12)

    fonction = emp.job_title.name if emp.job_title else "—"
    section = extra.get("section") or (emp.section.name if emp.section else (emp.department.name if emp.department else "—"))
    factory = emp.factory.name if emp.factory else "—"

    mergefields = {
        "Nom": emp.last_name or "—",
        "Prénom": emp.first_name or "—",
        "Matricule": emp.employee_id or "—",
        "Fonction": fonction,
        "Section": section,
        "Site_Publi": factory,
        "DE_en_LETTRES": hire_fr,
        "D_Fin_6_en_lettres": fin6,
        "D_Fin_12_en_lettres": fin12,
        "Début_1_ère_renouvellement1": fin6,
        "Fin_1ère_renouvellement1": fin12,
    }
    return fill_template(
        os.path.join(TEMPLATES_DIR, "cdd-12mois.docx"),
        {},
        mergefields=mergefields,
    )

def replace_mergefields(doc, field_map):
    """Remplace les champs MERGEFIELD Word (complexes ET simples) par du texte statique."""
    root = doc.element

    # ── 1. Champs simples : <w:fldSimple w:instr=" MERGEFIELD X "> ... </w:fldSimple>
    fld_simples = root.findall('.//' + qn('w:fldSimple'))
    for fs in fld_simples:
        instr = fs.get(qn('w:instr')) or ""
        m = re.search(r'MERGEFIELD\s+"?([^"\s]+)"?', instr)
        if not m:
            continue
        field_name = m.group(1)
        value = field_map.get(field_name, "")

        existing_run = fs.find(qn('w:r'))
        rpr = existing_run.find(qn('w:rPr')) if existing_run is not None else None

        new_r = OxmlElement('w:r')
        if rpr is not None:
            new_r.append(copy.deepcopy(rpr))
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = str(value)
        new_r.append(t)

        for child in list(fs):
            fs.remove(child)
        fs.append(new_r)

    # ── 2. Champs complexes : fldChar begin/instrText/separate/end ──────────
    paragraphs = root.findall('.//' + qn('w:p'))
    for p in paragraphs:
        runs = list(p.findall(qn('w:r')))
        i = 0
        while i < len(runs):
            r = runs[i]
            fld = r.find(qn('w:fldChar'))
            if fld is not None and fld.get(qn('w:fldCharType')) == 'begin':
                collected = [r]
                field_name = None
                j = i + 1
                while j < len(runs):
                    rj = runs[j]
                    collected.append(rj)
                    instr = rj.find(qn('w:instrText'))
                    if instr is not None and instr.text and 'MERGEFIELD' in instr.text:
                        m = re.search(r'MERGEFIELD\s+"?([^"\s]+)"?', instr.text)
                        if m:
                            field_name = m.group(1)
                    fld2 = rj.find(qn('w:fldChar'))
                    if fld2 is not None and fld2.get(qn('w:fldCharType')) == 'separate':
                        j += 1
                        break
                    j += 1
                template_run = None
                k = j
                while k < len(runs):
                    rk = runs[k]
                    collected.append(rk)
                    if template_run is None and rk.find(qn('w:t')) is not None:
                        template_run = rk
                    fldend = rk.find(qn('w:fldChar'))
                    if fldend is not None and fldend.get(qn('w:fldCharType')) == 'end':
                        break
                    k += 1

                value = field_map.get(field_name, "") if field_name else ""

                new_r = OxmlElement('w:r')
                if template_run is not None:
                    rpr = template_run.find(qn('w:rPr'))
                    if rpr is not None:
                        new_r.append(copy.deepcopy(rpr))
                t = OxmlElement('w:t')
                t.set(qn('xml:space'), 'preserve')
                t.text = str(value)
                new_r.append(t)

                collected[0].addprevious(new_r)
                for el in collected:
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)

                runs = list(p.findall(qn('w:r')))
                i = runs.index(new_r) + 1
                continue
            i += 1

def build_convocation_abandon(emp, extra, template_name):
    """Commun aux versions CDD et CDI de la lettre de convocation."""
    mergefields = {
        "Nom": emp.last_name or "\u2014",
        "Pr\u00e9nom": emp.first_name or "\u2014",
        "Matricule": emp.employee_id or "\u2014",
        "Emploi_occup\u00e9": emp.job_title.name if emp.job_title else "\u2014",
        "ADRESSE": getattr(emp, "address", None) or "\u2014",
    }
    reps = {
        "<date_now>"      : today_fr(),
        "<date_embauche>" : fmt_date(emp.hire_date),
        "<date_abandon>"  : fmt_date(extra.get("date_abandon")),
        "<date_presence>" : fmt_date(extra.get("date_derniere_presence")),
        "<chiffres>"      : extra.get("montant") or "\u2014",
    }
    return fill_template(
        os.path.join(TEMPLATES_DIR, template_name),
        reps, mergefields=mergefields,
    )

def build_convocation_cdd(emp, extra):
    return build_convocation_abandon(
        emp, extra, "LETTRE_DE_CONVOCATION_POUR_ABANDON_DE_POSTE_-_CDD.docx"
    )

def build_convocation_cdi(emp, extra):
    return build_convocation_abandon(
        emp, extra, "LETTRE_DE_CONVOCATION_POUR_ABANDON_DE_POSTE_-_CDI.docx"
    )

def build_suspension(emp, extra):
    fonction = emp.job_title.name if emp.job_title else "\u2014"
    reps = {
        "<date_jour>"          : today_fr(),
        "<nom_employe>"        : f"{emp.last_name} {emp.first_name}",
        "<poste_employe>"      : fonction,
        "<matricule>"          : emp.employee_id or "\u2014",
        "<ref_suspension>"     : extra.get("ref") or ref_rh(),
        "<date_certificat>"    : fmt_date(extra.get("date_certificat")),
        "<texte_medecin>"      : extra.get("delivre_par") or "\u2014",
        "<nombre_jours_arret>" : extra.get("jours_arret") or "\u2014",
        "<date_suspension>"    : fmt_date(extra.get("date_suspension")),
        "<nombre_preavis>"     : extra.get("duree_preavis") or "\u2014",
        "<date_preavis_debut>" : fmt_date(extra.get("date_preavis_debut")),
        "<date_preavis_fin>"   : fmt_date(extra.get("date_preavis_fin")),
    }
    return fill_template(
        os.path.join(TEMPLATES_DIR, "Mod\u00e8le_suspension_du_contrat.docx"),
        reps,
    )
# ── API ───────────────────────────────────────────────────────────────────────

# BUILDERS = {
#     "attestation"   : (build_attestation,    "Attestation_emploi"),
#     "certificat"    : (build_certificat,     "Certificat_travail"),
#     "contrat_cdd"   : (build_contrat_cdd,    "Contrat_CDD"),
#     "evaluation_cdd": (build_evaluation_cdd, "Evaluation_CDD"),
#     "confirmation_cdi" : (build_confirmation_cdi,  "Confirmation_CDI"),
# }
BUILDERS = {
    "attestation"     : (build_attestation,      "Attestation_emploi"),
    "certificat"      : (build_certificat,       "Certificat_travail"),
    "cdd_3"           : (build_cdd_3,            "Evaluation_CDD_3mois"),
    "cdd_6"           : (build_cdd_6,            "Contrat_CDD_6mois"),      # ex build_contrat_cdd, renommé
    "cdd_12"          : (build_cdd_12,           "Evaluation_CDD_12mois"),
    "cdd_18"          : (build_cdd_18,           "Contrat_CDD_18mois"),     # ex build_confirmation_cdi, renommé
    "evaluation_cdd"  : (build_evaluation_cdd,   "Evaluation_CDD"),
    "convocation_cdd"  : (build_convocation_cdd,   "Convocation_Abandon_Poste_CDD"),
    "convocation_cdi"  : (build_convocation_cdi,   "Convocation_Abandon_Poste_CDI"),
    "suspension"       : (build_suspension,        "Suspension_Contrat"),
    "badge"           : (None,                   "Badge_employe"), 
}

@api_view(["POST"])
@permission_classes([IsHRUser])
def generate_document(request, employee_id, doc_type):
    if doc_type not in BUILDERS:
        return Response({"detail": f"Type inconnu: {doc_type}"}, status=400)
    if doc_type == "badge":
        return Response({"detail": "Le badge n'est disponible qu'en PDF. Utilisez l'endpoint /pdf/."}, status=400)
    try:
        emp = Employee.objects.select_related(
            "factory",
            "department",
            "section",
            "classification",
            "job_title",
        ).get(pk=employee_id)
    except Employee.DoesNotExist:
        return Response({"detail": "Employ\u00e9 introuvable."}, status=404)
    fn, prefix = BUILDERS[doc_type]
    try:
        doc_bytes = fn(emp, request.data or {})
    except FileNotFoundError as e:
        return Response({"detail": f"Template manquant: {e}"}, status=500)
    except Exception as e:
        import traceback
        return Response({"detail": str(e), "trace": traceback.format_exc()}, status=500)
    safe = re.sub(r"[^\w\-.]","_",f"{prefix}_{emp.last_name}_{emp.first_name}.docx")
    resp = HttpResponse(doc_bytes,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    resp["Content-Disposition"] = f'attachment; filename="{safe}"'
    return resp

@api_view(["POST"])
@permission_classes([IsHRUser])
def bulk_documents_zip(request):
    doc_type = request.data.get("document_type")
    employee_ids = request.data.get("employee_ids", [])
    extra = request.data.get("extra", {}) or {}

    if doc_type not in BUILDERS:
        return Response({"detail": f"Type inconnu: {doc_type}"}, status=400)
    if not employee_ids:
        return Response({"detail": "Aucun employé sélectionné."}, status=400)

    employees = Employee.objects.select_related(
        "factory", "department", "classification", "job_title",
    ).filter(employee_id__in=employee_ids)

    if not employees.exists():
        return Response({"detail": "Aucun employé trouvé."}, status=404)

    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        if doc_type == "badge":
            for emp in employees:
                try:
                    pdf_bytes = build_badge_pdf(emp)
                    filename = re.sub(r"[^\w\-.]", "_", f"Badge_{emp.employee_id}_{emp.last_name}_{emp.first_name}.pdf")
                    zip_file.writestr(filename, pdf_bytes)
                except Exception as e:
                    error_filename = re.sub(r"[^\w\-.]", "_", f"ERREUR_{emp.employee_id}_{emp.last_name}_{emp.first_name}.txt")
                    zip_file.writestr(error_filename, str(e))
        else:
            fn, prefix = BUILDERS[doc_type]
            for emp in employees:
                try:
                    doc_bytes = fn(emp, extra)
                    filename = re.sub(r"[^\w\-.]", "_", f"{prefix}_{emp.employee_id}_{emp.last_name}_{emp.first_name}.docx")
                    zip_file.writestr(filename, doc_bytes)
                except Exception as e:
                    error_filename = re.sub(r"[^\w\-.]", "_", f"ERREUR_{emp.employee_id}_{emp.last_name}_{emp.first_name}.txt")
                    zip_file.writestr(error_filename, str(e))

    zip_buffer.seek(0)
    zip_name = re.sub(r"[^\w\-.]", "_", f"documents_rh_{doc_type}.zip")
    resp = HttpResponse(zip_buffer.getvalue(), content_type="application/zip")
    resp["Content-Disposition"] = f'attachment; filename="{zip_name}"'
    return resp


# @api_view(["GET"])
# @permission_classes([IsHRUser])
# def list_templates(request):
#     return Response({"templates":[
#         {"id":"attestation",    "title":"Attestation d\u2019emploi"},
#         {"id":"certificat",     "title":"Certificat de travail"},
#         {"id":"contrat_cdd",    "title":"Contrat CDD 6 mois"},
#         {"id":"evaluation_cdd", "title":"\u00c9valuation & Renouvellement CDD"},
#         {"id":"confirmation_cdi", "title":"Évaluation & Confirmation CDI"},
#     ]})
@api_view(["GET"])
@permission_classes([IsHRUser])
def list_templates(request):
    return Response({"templates":[
        {"id":"attestation",     "title":"Attestation d\u2019emploi"},
        {"id":"certificat",      "title":"Certificat de travail"},
        {"id":"cdd_3",           "title":"Contrat CDD 3 mois"},
        {"id":"cdd_6",           "title":"Contrat CDD 6 mois"},
        {"id":"cdd_12",          "title":"Contrat CDD 12 mois"},
        {"id":"cdd_18",          "title":"Contrat CDD 18 mois"},
        {"id":"evaluation_cdd",  "title":"\u00c9valuation & Renouvellement CDD"},
        {"id":"badge",           "title":"Badge employé"},
    ]})

@api_view(["POST"])
@permission_classes([IsHRUser])
def generate_document_pdf(request, employee_id, doc_type):
    if doc_type not in BUILDERS:
        return Response({"detail": f"Type inconnu: {doc_type}"}, status=400)
    try:
        emp = Employee.objects.select_related(
            "factory", "department", "section", "classification", "job_title",
        ).get(pk=employee_id)
    except Employee.DoesNotExist:
        return Response({"detail": "Employé introuvable."}, status=404)

    # ── Cas spécial : badge (HTML → PDF direct, pas de docx/LibreOffice) ──
    if doc_type == "badge":
        try:
            pdf_bytes = build_badge_pdf(emp)
        except Exception as e:
            import traceback
            return Response({"detail": str(e), "trace": traceback.format_exc()}, status=500)
        safe = re.sub(r"[^\w\-.]", "_", f"Badge_{emp.last_name}_{emp.first_name}.pdf")
        resp = HttpResponse(pdf_bytes, content_type="application/pdf")
        resp["Content-Disposition"] = f'inline; filename="{safe}"'
        return resp

    fn, prefix = BUILDERS[doc_type]
    try:
        doc_bytes = fn(emp, request.data or {})
    except FileNotFoundError as e:
        return Response({"detail": f"Template manquant: {e}"}, status=500)
    except Exception as e:
        import traceback
        return Response({"detail": str(e), "trace": traceback.format_exc()}, status=500)

    # Write docx to temp file and convert to PDF via LibreOffice
    import tempfile
    import subprocess
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "document.docx")
        with open(docx_path, "wb") as f:
            f.write(doc_bytes)
        try:
            subprocess.run(  # nosec B603 B607 - docx_path built server-side from tempfile + fixed filename, not user input
                ["libreoffice", "--headless", "--convert-to", "pdf",
                 "--outdir", tmpdir, docx_path],
                timeout=30, check=True,
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
        except subprocess.TimeoutExpired:
            return Response({"detail": "Conversion PDF timeout."}, status=500)
        except subprocess.CalledProcessError as e:
            return Response({"detail": f"Erreur conversion PDF: {e}"}, status=500)

        pdf_path = os.path.join(tmpdir, "document.pdf")
        if not os.path.exists(pdf_path):
            return Response({"detail": "PDF non généré."}, status=500)

        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()

    safe = re.sub(r"[^\w\-.]", "_", f"{prefix}_{emp.last_name}_{emp.first_name}.pdf")
    resp = HttpResponse(pdf_bytes, content_type="application/pdf")
    resp["Content-Disposition"] = f'inline; filename="{safe}"'
    return resp



@api_view(["POST"])
@permission_classes([IsHRUser])
def bulk_documents_pdf(request):
    import tempfile
    import subprocess
    doc_type     = request.data.get("document_type")
    employee_ids = request.data.get("employee_ids", [])
    extra        = request.data.get("extra", {}) or {}

    if doc_type not in BUILDERS:
        return Response({"detail": f"Type inconnu: {doc_type}"}, status=400)
    if not employee_ids:
        return Response({"detail": "Aucun employe selectionne."}, status=400)

    emps = Employee.objects.select_related(
        "factory", "department", "section", "classification", "job_title",
    ).filter(employee_id__in=employee_ids)

    if not emps.exists():
        return Response({"detail": "Aucun employe trouve."}, status=404)

    pdf_bytes_list = []

    if doc_type == "badge":
        for emp in emps:
            try:
                pdf_bytes_list.append(build_badge_pdf(emp))
            except Exception:
                logger.exception("bulk_documents_pdf: badge failed for employee %s", emp.employee_id)
                continue
    else:
        fn, prefix = BUILDERS[doc_type]
        with tempfile.TemporaryDirectory() as tmpdir:
            for emp in emps:
                try:
                    doc_bytes = fn(emp, extra)
                    docx_path = os.path.join(tmpdir, f"{emp.employee_id}.docx")
                    with open(docx_path, "wb") as f:
                        f.write(doc_bytes)
                    subprocess.run(
                        ["libreoffice", "--headless", "--convert-to", "pdf",
                         "--outdir", tmpdir, docx_path],
                        timeout=30, check=True,
                        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                    )
                    pdf_path = os.path.join(tmpdir, f"{emp.employee_id}.pdf")
                    if os.path.exists(pdf_path):
                        with open(pdf_path, "rb") as f:
                            pdf_bytes_list.append(f.read())
                    else:
                        logger.error("bulk_documents_pdf: pdf not produced for employee %s (docx_path=%s)", emp.employee_id, docx_path)
                except Exception:
                    logger.exception("bulk_documents_pdf: unexpected error for employee %s", emp.employee_id)
                    continue

    if not pdf_bytes_list:
        return Response({"detail": "Aucun PDF genere."}, status=500)

    # Merge PDFs using pypdf (inchangé)
    import tempfile as _tempfile
    with _tempfile.TemporaryDirectory() as tmpdir:
        merged_path = os.path.join(tmpdir, "merged.pdf")
        try:
            from pypdf import PdfWriter
            writer = PdfWriter()
            for pdf_bytes in pdf_bytes_list:
                import io as _io
                writer.append(_io.BytesIO(pdf_bytes))
            with open(merged_path, "wb") as f:
                writer.write(f)
            with open(merged_path, "rb") as f:
                final_bytes = f.read()
        except Exception:
            final_bytes = pdf_bytes_list[0]

    resp = HttpResponse(final_bytes, content_type="application/pdf")
    resp["Content-Disposition"] = f'inline; filename="documents_rh_{doc_type}.pdf"'
    return resp